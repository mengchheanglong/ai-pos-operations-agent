"""
Generate a professional, fully styled Word Document (.docx) report
strictly following the 9 required sections from the assignment guidelines.
"""
import os
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(PROJECT_ROOT, "docs")
SCREENSHOTS_DIR = os.path.join(PROJECT_ROOT, "screenshots")
OUTPUT_DOCX = os.path.join(DOCS_DIR, "AI_Business_Operations_Agent_Report.docx")

# Theme Colors
C_PRIMARY = RGBColor(30, 64, 175)     # #1e40af (Royal Blue)
C_SECONDARY = RGBColor(55, 65, 81)   # #374151 (Dark Slate)
C_MUTED = RGBColor(100, 116, 139)    # #64748b (Muted Gray)
C_DARK = RGBColor(15, 23, 42)        # #0f172a (Charcoal Black)
C_GREEN = RGBColor(5, 150, 105)      # #059669 (Emerald)
C_AMBER = RGBColor(217, 119, 6)      # #d97706 (Amber)


def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Set thin clean borders on a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def format_table_header(row, col_widths, hex_color="1E40AF"):
    """Format table header row with background and white bold text."""
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]
        set_cell_background(cell, hex_color)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)


def add_callout(doc, text, title=None, border_hex="2563EB", bg_hex="EFF6FF"):
    """Add a styled callout box using a 1x1 table."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    
    # Left border only
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    if title:
        r_title = p.add_run(f"{title}\n")
        r_title.bold = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = RGBColor(30, 64, 175)
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(doc, code_str):
    """Add a monospace formatted code box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(code_str)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_word_report():
    doc = Document()

    # Page Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = C_DARK

    # ==========================================
    # HEADER / TITLE
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("AURA POS — AI Smart Retail & Operations Terminal")
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = C_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Autonomous Conversational Checkout, Multi-Path AI Decision Routing, and Automated Operations\nAssignment Submission Report")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = C_MUTED

    # Meta Info Card
    meta_tbl = doc.add_table(rows=1, cols=3)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cells = meta_tbl.rows[0].cells
    meta_cells[0].width = Inches(2.1)
    meta_cells[1].width = Inches(2.1)
    meta_cells[2].width = Inches(2.3)
    for c in meta_cells:
        set_cell_background(c, "F1F5F9")
    meta_cells[0].paragraphs[0].add_run("Candidate: ").bold = True
    meta_cells[0].paragraphs[0].add_run("Mengchheang Long")
    meta_cells[1].paragraphs[0].add_run("Model: ").bold = True
    meta_cells[1].paragraphs[0].add_run("DeepSeek-V3 LLM")
    meta_cells[2].paragraphs[0].add_run("Framework: ").bold = True
    meta_cells[2].paragraphs[0].add_run("Python Agent / n8n")
    for c in meta_cells:
        for p in c.paragraphs:
            p.runs[0].font.size = Pt(9)
            p.runs[1].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================================
    # 1. PROBLEM STATEMENT
    # ==========================================
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Problem Statement")
    r1.font.color.rgb = C_PRIMARY
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Small-to-medium retail shops, boutique stores, and counter sales operations frequently struggle with managing customer transactions, "
        "shelf stock level verification, customer inquiries, and order approval governance through disjointed manual workflows. "
        "Store cashiers and staff typically ring up or take phone orders manually, cross-reference spreadsheet inventory by hand, "
        "manually check whether high-value purchases require manager override approval, and struggle to coordinate out-of-stock color variants on the fly."
    )

    doc.add_paragraph(
        "This manual mode of counter and sales operation introduces severe operational friction:"
    )

    p_pts = doc.add_paragraph()
    p_pts.paragraph_format.left_indent = Inches(0.25)
    p_pts.paragraph_format.space_after = Pt(4)
    p_pts.add_run("• Slow Checkout & Phone Order Ring-Up: ").bold = True
    p_pts.add_run("Cashiers spend precious minutes manually deciphering customer requests, typing SKUs, and checking shelf counts, causing customer queues at the counter.\n")
    p_pts.add_run("• Stock Inaccuracies & Shelf Desynchronization: ").bold = True
    p_pts.add_run("Discrepancies between physical shelf stock and central inventory sheets cause stockouts and overselling during peak hours.\n")
    p_pts.add_run("• Inconsistent Manager Approval Governance: ").bold = True
    p_pts.add_run("Staff members apply discount rules, high-ticket escalation thresholds, and substitution criteria inconsistently across shifts.\n")
    p_pts.add_run("• High Friction on Product Substitutions: ").bold = True
    p_pts.add_run("When a customer requests an out-of-stock color or size, staff rarely have instant visibility into backstock alternatives, leading to abandoned sales.")

    add_callout(
        doc,
        "Target Real-World Problem: Manual retail point-of-sale and counter operations suffer from delayed order fulfillment, stock desynchronization, "
        "and human error. An intelligent autonomous AI POS terminal is needed to bridge natural customer communications, live shelf inventory, "
        "business rules (e.g. Manager Override for >= $100), and instant multi-party alerts with zero manual data entry.",
        title="Problem Definition Summary"
    )

    # ==========================================
    # 2. PROPOSED AI AGENT SOLUTION
    # ==========================================
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Proposed AI Agent Solution")
    r2.font.color.rgb = C_PRIMARY
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "We propose AURA POS — an autonomous end-to-end AI Smart Retail and Order Operations Terminal designed to "
        "receive conversational counter, phone, or messaging order requests, parse unstructured intent using the DeepSeek Large Language Model (deepseek-chat), "
        "query real-time shelf stock from Google Sheets, autonomously evaluate fulfillment and governance rules (e.g., auto-approving small sales vs. escalating high-value orders for manager override), "
        "and route the workflow through five distinct operational paths while sending real-time Telegram alerts and receipts."
    )

    doc.add_paragraph(
        "Key Pillars of the Solution:"
    )
    p_sol = doc.add_paragraph()
    p_sol.paragraph_format.left_indent = Inches(0.25)
    p_sol.add_run("1. Conversational POS Checkout: ").bold = True
    p_sol.add_run("Understands natural, unstructured customer phrasing (e.g., 'Customer wants 2 black basic shirts' or '12 shirts for team event') without clunky catalog lookups.\n")
    p_sol.add_run("2. Real-Time Tool Integration: ").bold = True
    p_sol.add_run("Integrates live Google Sheets API for shelf inventory deduction, transaction logging, and register audit tracking, alongside Telegram Bot API for instant customer receipts and owner alerts.\n")
    p_sol.add_run("3. Multi-Path Autonomous Routing: ").bold = True
    p_sol.add_run("Autonomously determines whether to auto-approve (<$100), escalate for Manager Override (>= $100), recommend alternative variants on stock-outs, clarify missing details, or reject.\n")
    p_sol.add_run("4. Hybrid Terminal Orchestration: ").bold = True
    p_sol.add_run("Can operate as an interactive POS Web Terminal (Flask UI), standalone Python CLI terminal, or enterprise n8n workflow graph with resilient zero-downtime offline fallback.")

    # ==========================================
    # 3. WORKFLOW DIAGRAM
    # ==========================================
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. Workflow Diagram")
    r3.font.color.rgb = C_PRIMARY
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The operations workflow comprises 10 distinct operational capabilities organized into a 5-stage autonomous pipeline "
        "(surpassing the assignment minimum requirement of 5 steps). Stages 1 through 4 handle order ingestion, LLM intent extraction, "
        "and inventory reasoning. Stage 5 acts as an autonomous multi-path decision router that dispatches the request into one of five "
        "mutually exclusive operational branches (Approve, Escalate, Suggest Alternative, Clarify, Reject), executing branch-specific downstream "
        "actions (stock deduction, order logging, Telegram alerts, and audit logging) before returning a structured receipt."
    )

    # Embed Workflow Diagram Image
    wf_img = os.path.join(SCREENSHOTS_DIR, "00_workflow_diagram.png")
    if os.path.exists(wf_img):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(wf_img, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run("Figure 1: End-to-End Autonomous AI Operations Architecture & Branch-Specific Routing")
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = C_MUTED

    # Workflow Steps Table
    wf_table_data = [
        ("Capability #", "Operational Stage", "Component / Tool", "Description & Execution Condition"),
        ("1", "Order Intake", "Webhook / Web POS / CLI", "Customer order enters system via Web POS terminal, Telegram, or webhook."),
        ("2", "Extract Order Intent", "DeepSeek LLM (Decision 1)", "AI extracts product name, variant, requested quantity, and intent into JSON."),
        ("3", "Read Live Inventory", "Google Sheets / Excel", "Fetches current shelf stock levels, SKU mappings, and threshold values."),
        ("4", "Evaluate Fulfillment", "DeepSeek LLM (Decisions 2 & 3)", "AI evaluates stock availability, threshold rules ($100 rule), and product validity."),
        ("5", "Autonomous Route", "5-Way AI Router", "Dynamically routes to 1 of 5 paths: approve, escalate, suggest_alt, clarify, reject."),
        ("6", "Deduct Stock", "Google Sheets / Excel", "Executed on APPROVE only: decrements physical shelf count in real time."),
        ("7", "Record Order", "Orders Database", "Executed on APPROVE: records order ID, customer, total price, and status."),
        ("8", "Customer Notification", "Telegram Bot API", "Transmits order confirmation (on approve), alternatives, or clarification prompt."),
        ("9", "Manager Escalation", "Telegram Bot API", "Executed on ESCALATE only: sends urgent override alert to owner for orders >= $100."),
        ("10", "Append Audit Trail", "Logs Database", "Executed on ALL paths: records immutable timestamped audit log of all decisions."),
    ]

    wf_tbl = doc.add_table(rows=len(wf_table_data), cols=4)
    wf_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(wf_tbl)
    col_w = [Inches(0.6), Inches(1.5), Inches(1.8), Inches(2.6)]

    for row_idx, data in enumerate(wf_table_data):
        row = wf_tbl.rows[row_idx]
        if row_idx == 0:
            for c_idx, text in enumerate(data):
                row.cells[c_idx].paragraphs[0].text = text
            format_table_header(row, col_w)
        else:
            for c_idx, text in enumerate(data):
                cell = row.cells[c_idx]
                cell.width = col_w[c_idx]
                p = cell.paragraphs[0]
                p.text = text
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(8.5)
            if row_idx % 2 == 1:
                for c in row.cells:
                    set_cell_background(c, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # 4. SYSTEM ARCHITECTURE
    # ==========================================
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. System Architecture")
    r4.font.color.rgb = C_PRIMARY
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The system follows a modular, decoupled architecture where the central Python Agent Core (agent.py) orchestrates "
        "data flow between external LLM reasoning APIs, live spreadsheet databases, and instant messaging bots. "
        "An alternative n8n workflow engine orchestration is also provided via importable JSON workflow templates."
    )

    # Embed Architecture Diagram Image
    arch_img = os.path.join(SCREENSHOTS_DIR, "00_architecture_diagram.png")
    if os.path.exists(arch_img):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(arch_img, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run("Figure 2: Modular System Architecture & Service Integration Diagram")
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = C_MUTED

    # Architecture Breakdown Table
    arch_table_data = [
        ("Tier / Component", "Technology Stack", "Role & Functional Responsibility"),
        ("Client Entry Points", "AURA POS Web Terminal / CLI / Webhook", "Accepts natural language orders from register screen, phone order desk, or API."),
        ("Agent Core & Router", "Python 3.10+ (agent.py)", "Coordinates 10-step pipeline, maintains state, applies business rules, and routes."),
        ("Reasoning & Decision Tier", "DeepSeek API (deepseek-chat)", "OpenAI-compatible LLM endpoint executing intent extraction and fulfillment logic."),
        ("Data Persistence Tier", "Google Sheets API (gspread)", "Multi-sheet relational database (Products, Orders, Customers, Logs)."),
        ("Resilient Fallback", "In-Memory Mock Database", "Provides zero-downtime offline execution if Google Cloud credentials are unavailable."),
        ("Notification Tier", "Telegram Bot API (@business0psBot)", "Dispatches instant digital receipts to customer and manager escalation alerts to owner."),
    ]

    arch_tbl = doc.add_table(rows=len(arch_table_data), cols=3)
    arch_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(arch_tbl)
    col_w_arch = [Inches(1.5), Inches(1.8), Inches(3.2)]

    for row_idx, data in enumerate(arch_table_data):
        row = arch_tbl.rows[row_idx]
        if row_idx == 0:
            for c_idx, text in enumerate(data):
                row.cells[c_idx].paragraphs[0].text = text
            format_table_header(row, col_w_arch)
        else:
            for c_idx, text in enumerate(data):
                cell = row.cells[c_idx]
                cell.width = col_w_arch[c_idx]
                p = cell.paragraphs[0]
                p.text = text
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(8.5)
            if row_idx % 2 == 1:
                for c in row.cells:
                    set_cell_background(c, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # 5. AI TOOLS, MODELS, AND EXTERNAL SERVICES USED
    # ==========================================
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. AI Tools, Models, and External Services Used")
    r5.font.color.rgb = C_PRIMARY
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "To satisfy the assignment requirement of integrating at least two external services, the agent integrates "
        "four external tools and services, coordinated by an LLM reasoning engine:"
    )

    tools_data = [
        ("Service Name", "Category", "Model / Protocol", "Detailed Role in Solution"),
        ("DeepSeek API", "LLM Reasoning", "deepseek-chat (V3)", "Performs structured intent extraction and decision reasoning using temperature=0.1."),
        ("Google Sheets API", "Database / Storage", "REST / gspread / OAuth2", "Provides persistent tables for Products shelf inventory, Orders, Customers, and Logs."),
        ("Telegram Bot API", "Messaging / Alerts", "HTTP Bot API (@business0psBot)", "Delivers real-time customer receipts and manager override alerts to owner chat."),
        ("n8n Orchestration", "Workflow Engine", "v2.8.4 Node Graph", "Visual node-based workflow alternative for enterprise retail automation environments."),
        ("Flask Web Framework", "POS Web Terminal", "WSGI / HTML5 / JS", "AURA POS cashier terminal interface featuring one-click presets, shelf stock, and live status."),
    ]

    tools_tbl = doc.add_table(rows=len(tools_data), cols=4)
    tools_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tools_tbl)
    col_w_t = [Inches(1.3), Inches(1.1), Inches(1.6), Inches(2.5)]

    for row_idx, data in enumerate(tools_data):
        row = tools_tbl.rows[row_idx]
        if row_idx == 0:
            for c_idx, text in enumerate(data):
                row.cells[c_idx].paragraphs[0].text = text
            format_table_header(row, col_w_t)
        else:
            for c_idx, text in enumerate(data):
                cell = row.cells[c_idx]
                cell.width = col_w_t[c_idx]
                p = cell.paragraphs[0]
                p.text = text
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(8.5)
            if row_idx % 2 == 1:
                for c in row.cells:
                    set_cell_background(c, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Detailed Sub-sections for Decision Points & Business Rules
    p_dp = doc.add_paragraph()
    p_dp.add_run("AI Decision Points & Autonomous Branching Details:\n").bold = True
    p_dp.add_run(
        "• Decision Point 1 (Intent Parsing): LLM extracts product, variant, quantity, and intent with confidence scoring.\n"
        "• Decision Point 2 (Fulfillment Evaluation): LLM checks requested quantity against inventory stock.\n"
        "• Decision Point 3 (Action Selection): LLM evaluates business rules to autonomously trigger 1 of 5 distinct actions:"
    )

    # Business Rules Table
    rules_data = [
        ("Rule ID", "Rule Name", "Condition Logic", "Autonomous Action", "Priority"),
        ("R001", "Auto-approve low value", "total_value < $100 AND stock >= quantity", "approve (Fast POS Sale)", "Priority 1"),
        ("R002", "Manager Override", "total_value >= $100 AND stock sufficient", "escalate (Manager Alert)", "Priority 2"),
        ("R003", "Suggest alternative", "stock < quantity AND alternative variant exists", "suggest_alternative (Shelf Stock-Out)", "Priority 3"),
        ("R004", "Reject unavailable", "stock < quantity AND no alternative exists", "reject (Unfulfillable)", "Priority 4"),
        ("R005", "Clarify missing info", "product not found OR quantity ambiguous", "clarify (Catalog Inquiry)", "Priority 5"),
    ]

    rules_tbl = doc.add_table(rows=len(rules_data), cols=5)
    rules_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(rules_tbl)
    col_w_r = [Inches(0.8), Inches(1.5), Inches(2.2), Inches(1.2), Inches(0.8)]

    for row_idx, data in enumerate(rules_data):
        row = rules_tbl.rows[row_idx]
        if row_idx == 0:
            for c_idx, text in enumerate(data):
                row.cells[c_idx].paragraphs[0].text = text
            format_table_header(row, col_w_r, hex_color="374151")
        else:
            for c_idx, text in enumerate(data):
                cell = row.cells[c_idx]
                cell.width = col_w_r[c_idx]
                p = cell.paragraphs[0]
                p.text = text
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(8.5)
            if row_idx % 2 == 1:
                for c in row.cells:
                    set_cell_background(c, "F8FAFC")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # 6. IMPLEMENTATION SCREENSHOTS
    # ==========================================
    h6 = doc.add_heading(level=1)
    r6 = h6.add_run("6. Implementation Screenshots")
    r6.font.color.rgb = C_PRIMARY
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Below are screenshots capturing the live execution of the AI Operations Agent across multiple scenarios, "
        "illustrating system health status, order entry, AI reasoning, and multi-path autonomous routing:"
    )

    screenshots_to_add = [
        ("01_web_ui_dashboard.png", "Figure 3: AURA POS Smart Terminal Dashboard with live DeepSeek, Telegram, and Sheets health indicators"),
        ("02_scenario_auto_approve.png", "Figure 4: Scenario 1 — Fast POS Checkout (Auto-Approval under $100 with immediate shelf stock deduction)"),
        ("03_scenario_escalate.png", "Figure 5: Scenario 2 — Manager Override Escalation to Store Owner for high-value sale ($120 >= $100)"),
        ("04_scenario_alternative.png", "Figure 6: Scenario 3 — Autonomous Variant Suggestion triggered when shelf stock is insufficient"),
        ("05_scenario_clarify.png", "Figure 7: Scenario 4 — Clarification Request triggered when an unrecognized item is requested at the register"),
    ]

    for img_name, caption in screenshots_to_add:
        img_path = os.path.join(SCREENSHOTS_DIR, img_name)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.space_before = Pt(6)
            r_img = p_img.add_run()
            r_img.add_picture(img_path, width=Inches(5.6))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run(caption)
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = C_MUTED

    # Note for user to add Telegram / Sheets screenshots if desired
    add_callout(
        doc,
        "Tip for Submission: You can also paste screenshots of your phone/Telegram chat receiving the alerts from @business0psBot "
        "or your Google Sheets spreadsheet right here in this Word document before exporting to PDF.",
        title="Customization Opportunity",
        border_hex="059669",
        bg_hex="ECFDF5"
    )

    # ==========================================
    # 7. SAMPLE EXECUTION (INPUT AND OUTPUT)
    # ==========================================
    h7 = doc.add_heading(level=1)
    r7 = h7.add_run("7. Sample Execution (Input and Output)")
    r7.font.color.rgb = C_PRIMARY
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The AI Agent was evaluated across four comprehensive end-to-end customer scenarios. Below are the exact "
        "inputs, step-by-step AI reasoning stages, and generated outputs:"
    )

    # Scenario 1
    doc.add_heading(level=2).add_run("Scenario 1: Fast Sale Checkout (Auto-Approve Flow)")
    doc.add_paragraph("Customer: John Doe | Input Message: \"I want 2 black Basic Shirts\"")
    add_code_block(doc,
"""Step 1 (Intent Extraction):
  Product: Basic Shirts | Variant: black | Quantity: 2 | Intent: order | Confidence: 0.95
Step 2 (Inventory Check):
  Found: Basic Shirt (Black) | Stock: 20 | Price: $10
Step 3 (AI Action Decision):
  Action: "approve"
  Reason: Total value is $20 (< $100) and stock (20) is sufficient for requested quantity (2).
Step 4 (Execution & Output):
  Stock updated: 20 -> 18 | Order recorded: ORD-B94D74AD | Customer Telegram Sent
Output JSON: {"status": "approved", "order_id": "ORD-B94D74AD", "total": 20}""")

    # Scenario 2
    doc.add_heading(level=2).add_run("Scenario 2: High-Value Counter Order (Manager Override Escalation Flow)")
    doc.add_paragraph("Customer: Jane Smith | Input Message: \"I need 12 black Basic Shirts\"")
    add_code_block(doc,
"""Step 1 (Intent Extraction):
  Product: Basic Shirts | Variant: black | Quantity: 12 | Intent: order | Confidence: 0.95
Step 2 (Inventory Check):
  Found: Basic Shirt (Black) | Stock: 18 | Price: $10
Step 3 (AI Action Decision):
  Action: "escalate"
  Reason: Total order value is $120, which is >= $100, and stock is sufficient (18 >= 12).
          Per business rules, orders with total value >= $100 require owner approval.
Step 4 (Execution & Output):
  Owner notified via Telegram alert with order breakdown | Status set to pending
Output JSON: {"status": "escalated", "order_id": "ORD-1434B6E7", "total": 120}""")

    # Scenario 3
    doc.add_heading(level=2).add_run("Scenario 3: Insufficient Shelf Stock (Autonomous Variant Suggestion Flow)")
    doc.add_paragraph("Customer: Bob Wilson | Input Message: \"I want 25 black Basic Shirts\"")
    add_code_block(doc,
"""Step 1 (Intent Extraction):
  Product: Basic Shirts | Variant: black | Quantity: 25 | Intent: order | Confidence: 0.95
Step 2 (Inventory Check):
  Found: Basic Shirt (Black) | Stock: 18 | Price: $10 (Insufficient for 25 units)
Step 3 (AI Action Decision):
  Action: "suggest_alternative"
  Reason: Requested quantity (25) exceeds available stock (18) for Basic Shirt in Black.
          Alternative variants available: White (18 in stock) and Red (15 in stock).
Step 4 (Execution & Output):
  Proposes breakdown: 18 Black Basic Shirts + 7 White Basic Shirts ($250 total)
  Customer notified via Telegram with confirmation prompt
Output JSON: {"status": "alternative_suggested", "order_id": "ORD-02692FB9", "total": 250}""")

    # Scenario 4
    doc.add_heading(level=2).add_run("Scenario 4: Unrecognized Item / SKU (Clarification Flow)")
    doc.add_paragraph("Customer: Alice Brown | Input Message: \"Do you have blue jeans?\"")
    add_code_block(doc,
"""Step 1 (Intent Extraction):
  Product: jeans | Variant: blue | Quantity: None | Intent: inquiry | Confidence: 0.90
Step 2 (Inventory Check):
  Query 'jeans' not matched in inventory catalogue.
Step 3 (AI Action Decision):
  Action: "clarify"
  Reason: Product 'blue jeans' was not found in active inventory.
Step 4 (Execution & Output):
  Polite message sent to customer requesting SKU clarification or alternative inquiry.
Output JSON: {"status": "clarify", "order_id": "ORD-82D6DABA"}""")

    # ==========================================
    # 8. VIDEO DEMO LINK
    # ==========================================
    h8 = doc.add_heading(level=1)
    r8 = h8.add_run("8. Video Demo Link")
    r8.font.color.rgb = C_PRIMARY
    h8.paragraph_format.space_before = Pt(14)
    h8.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "A comprehensive video walkthrough demonstrating the live agent processing order requests, triggering AI decisions, "
        "and sending real-time Telegram alerts has been recorded and uploaded for evaluation:"
    )

    add_callout(
        doc,
        "Public Video URL: [PASTE YOUR PUBLIC VIDEO LINK HERE - e.g., YouTube Unlisted, Loom, or Google Drive URL]\n\n"
        "Demo Content Covered in Video:\n"
        "1. Overview of the retail POS problem and autonomous architecture.\n"
        "2. Live transaction ring-up in AURA POS Web Terminal (http://localhost:5000) and CLI (python agent.py).\n"
        "3. Live Telegram bot (@business0psBot) alerts popping up for Customer digital receipts and Manager Escalation approvals.\n"
        "4. Demonstration of autonomous 5-way branching: Fast Sale (<$100), Manager Override (>=$100), Alternative suggestion, and Clarification.\n"
        "5. Inspection of Google Sheets records (Products shelf counts, Orders register, Customers, and audit Logs).",
        title="Public Video Demo Link (Clickable)",
        border_hex="D97706",
        bg_hex="FFFBEB"
    )

    # ==========================================
    # 9. LINK TO YOUR PROJECT
    # ==========================================
    h9 = doc.add_heading(level=1)
    r9 = h9.add_run("9. Link to Your Project")
    r9.font.color.rgb = C_PRIMARY
    h9.paragraph_format.space_before = Pt(14)
    h9.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The project source code, workflows, schemas, and demo suite are publicly accessible on GitHub:"
    )

    add_callout(
        doc,
        "GitHub Public Repository:\nhttps://github.com/mengchheanglong/business-operations-agent\n\n"
        "Repository Structure:\n"
        "• agent.py - Main autonomous AI operations agent with DeepSeek & Telegram integration\n"
        "• web_app.py & templates/index.html - Interactive AURA POS Smart Terminal web application\n"
        "• demo.py & demo/demo-results.json - Multi-path test suite and recorded outputs\n"
        "• workflows/order-agent-workflow.json - 17-node n8n orchestration workflow template\n"
        "• workflows/setup-sheets-workflow.json - Google Sheets automated initialization template\n"
        "• docs/ - Full report, credentials guide, and technical specs\n"
        "• requirements.txt & .env.example - Dependency management and environment templates",
        title="Project Link & Source Code",
        border_hex="1E40AF",
        bg_hex="EFF6FF"
    )

    doc.add_paragraph("Quick Start Instructions for Evaluators:").bold = True
    add_code_block(doc,
"""# 1. Clone repository
git clone https://github.com/mengchheanglong/business-operations-agent.git
cd business-operations-agent

# 2. Install dependencies
pip install -r requirements.txt

# 3. Configure API keys (DeepSeek, Telegram, Google Sheets in .env)
cp .env.example .env

# 4. Run interactive web dashboard
python web_app.py
# Open http://localhost:5000 in your browser

# 5. Or execute automated demo suite across all 4 scenarios
python demo.py""")

    # ==========================================
    # APPENDIX: REQUIREMENTS VERIFICATION
    # ==========================================
    h_app = doc.add_heading(level=1)
    r_app = h_app.add_run("10. Requirements Verification Matrix")
    r_app.font.color.rgb = C_PRIMARY
    h_app.paragraph_format.space_before = Pt(14)
    h_app.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "The table below verifies full compliance with all criteria from the assignment specification:"
    )

    matrix_data = [
        ("Criterion", "Assignment Minimum", "Implemented in This Project", "Compliance"),
        ("Problem Scope", "Clearly defined real-world problem", "Retail POS terminal & small business inventory operations", "MET (100%)"),
        ("LLM Reasoning", "Required for decisions", "DeepSeek LLM (deepseek-chat) for intent & action reasoning", "MET (100%)"),
        ("External Tools", ">= 2 external tools/services", "4 services: Google Sheets, Telegram Bot, DeepSeek, n8n", "EXCEEDED (4 tools)"),
        ("Workflow Steps", ">= 5 steps", "10 sequential operational steps", "EXCEEDED (10 steps)"),
        ("AI Decision Points", ">= 2 decision points", "3 decision points (Intent, Stock fulfillment, Action routing)", "EXCEEDED (3 decisions)"),
        ("Autonomous Decision", ">= 1 workflow-changing decision", "5-way autonomous branching (approve, escalate, suggest, clarify, reject)", "EXCEEDED (5 paths)"),
        ("Deliverables Format", "One PDF with sections 1 to 9", "All 9 sections structured with embedded diagrams & screenshots", "READY"),
    ]

    mat_tbl = doc.add_table(rows=len(matrix_data), cols=4)
    mat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(mat_tbl)
    col_w_m = [Inches(1.3), Inches(1.5), Inches(2.7), Inches(1.0)]

    for row_idx, data in enumerate(matrix_data):
        row = mat_tbl.rows[row_idx]
        if row_idx == 0:
            for c_idx, text in enumerate(data):
                row.cells[c_idx].paragraphs[0].text = text
            format_table_header(row, col_w_m)
        else:
            for c_idx, text in enumerate(data):
                cell = row.cells[c_idx]
                cell.width = col_w_m[c_idx]
                p = cell.paragraphs[0]
                p.text = text
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.runs[0].font.size = Pt(8.5)
            if row_idx % 2 == 1:
                for c in row.cells:
                    set_cell_background(c, "F8FAFC")

    doc.save(OUTPUT_DOCX)
    print(f"Word report successfully generated at: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_word_report()
